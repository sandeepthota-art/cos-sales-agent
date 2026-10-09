"""Fills Actual Label for the 3 blind-test rows (ambiguity rule validation)
and the recovered "Re: Our call last week" row, from the real run results.

Run once: python eval/fill_blind_test_results.py
"""

from openpyxl import load_workbook

RESULTS = {
    "Thoughts?": "1. Undecided",
    "Thoughts on the attached pricing proposal?": "1. Needs reply",
    "Can you approve this by 3 PM?": "1. Needs reply: ASAP",
    "Re: Our call last week": "1. Needs reply",
}

wb = load_workbook("eval/label_eval_set.xlsx")
ws = wb["Eval Set"]

filled = []
for row in range(2, ws.max_row + 1):
    subject = ws.cell(row=row, column=5).value
    if subject in RESULTS:
        ws.cell(row=row, column=10).value = RESULTS[subject]
        filled.append(subject)

wb.save("eval/label_eval_set.xlsx")
print(f"Filled {len(filled)} rows: {filled}")
