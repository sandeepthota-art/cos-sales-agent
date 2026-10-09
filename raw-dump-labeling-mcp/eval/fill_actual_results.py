"""Fills eval/label_eval_set.xlsx with the real-world test run results:
actual sender address used + actual label assigned, matched to rows by
Subject. Rows with no result (not found/not sent) are left untouched.

Run once: python eval/fill_actual_results.py
"""

from openpyxl import load_workbook

# subject -> (real sender address, actual label)
RESULTS = {
    "PO details needed today to lock in pricing": ("deepanshu.sonwane@databeat.io", "1. Needs reply: ASAP"),
    "Thoughts?": ("deepanshu.sonwane@databeat.io", "1. Needs reply"),
    "Circling back on the deck": ("deepanshu.sonwane@mediamint.com", "1. Needs reply"),
    "Need confirmation -- Monday EOD deliverable": ("deepanshu.sonwane@mediamint.com", "1. Needs reply: ASAP"),
    "Fwd: Pricing doc": ("deepanshu.sonwane@mediamint.com", "1. Needs reply"),
    "Just checking in on our proposal": ("ironman.primus@gmail.com", "1. Needs reply"),
    "20% off if you sign this month": ("ironman.primus@gmail.com", "1. Needs reply"),
    "Meeting Response: Declined": ("ironman.primus@gmail.com", "1. Delete"),
    "Renewal terms -- contract expires in 3 days": ("deepanshusonwane1@gmail.com", "1. Needs reply: ASAP"),
    "Signed contract attached": ("deepanshusonwane1@gmail.com", "1. Needs reply: ASAP"),
}

wb = load_workbook("eval/label_eval_set.xlsx")
ws = wb["Eval Set"]

matched, mismatched, not_found = [], [], []
for row in range(2, ws.max_row + 1):
    subject = ws.cell(row=row, column=5).value
    expected = ws.cell(row=row, column=7).value
    if subject not in RESULTS:
        not_found.append(subject)
        continue
    real_sender, actual = RESULTS[subject]
    ws.cell(row=row, column=2).value = real_sender  # Sender Email
    ws.cell(row=row, column=10).value = actual       # Actual Label
    if expected == actual:
        matched.append(subject)
    else:
        mismatched.append((subject, expected, actual))

wb.save("eval/label_eval_set.xlsx")

print(f"Filled {len(matched) + len(mismatched)} rows.")
print(f"MATCH: {len(matched)}")
for s in matched:
    print(f"   {s}")
print(f"MISMATCH: {len(mismatched)}")
for s, exp, act in mismatched:
    print(f"   {s} -- expected {exp!r}, got {act!r}")
print(f"Not found in this run: {len(not_found)}")
for s in not_found:
    print(f"   {s}")
