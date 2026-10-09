"""Fills Gmail Message ID (and corrects placeholder sender addresses on the
3 blind-test rows to the real geetasonwane6@gmail.com) from the final
14-email results table.

Run once: python eval/fill_gmail_ids.py
"""

from openpyxl import load_workbook

# subject -> (gmail_message_id, real_sender_or_None)
RESULTS = {
    "Re: Our call last week": ("1a11f273a774ec6b", None),
    "PO details needed today to lock in pricing": ("1a11f27f619f0ae7", None),
    "Circling back on the deck": ("1a11f2997e38410f", None),
    "Need confirmation -- Monday EOD deliverable": ("1a11f2a18b37359a", None),
    "Fwd: Pricing doc": ("1a11f2a9a6a5a11a", None),
    "Just checking in on our proposal": ("1a11f2b4ec5ddd49", None),
    "20% off if you sign this month": ("1a11f2bb49e26021", None),
    "Meeting Response: Declined": ("1a11f2c1ffad09e4", None),
    "Renewal terms -- contract expires in 3 days": ("1a11f2cd7aebd3a0", None),
    "Signed contract attached": ("1a11f2d60eceb143", None),
    "Thoughts on the attached pricing proposal?": ("1a11f6095ce1b888", "geetasonwane6@gmail.com"),
    "Can you approve this by 3 PM?": ("1a11f611792226ac", "geetasonwane6@gmail.com"),
}

wb = load_workbook("eval/label_eval_set.xlsx")
ws = wb["Eval Set"]

filled = []
for row in range(2, ws.max_row + 1):
    subject = ws.cell(row=row, column=5).value
    sender = ws.cell(row=row, column=2).value

    if subject == "Thoughts?":
        # Two rows share this subject -- disambiguate by current sender.
        if sender == "deepanshu.sonwane@databeat.io":
            ws.cell(row=row, column=9).value = "1a11f28a275ea427"
            filled.append(f"row {row} (databeat.io)")
        elif sender == "ambiguity-test6@example.com":
            ws.cell(row=row, column=9).value = "1a11f5fc7a55394c"
            ws.cell(row=row, column=2).value = "geetasonwane6@gmail.com"
            filled.append(f"row {row} (geetasonwane6)")
        continue

    if subject in RESULTS:
        gmail_id, real_sender = RESULTS[subject]
        ws.cell(row=row, column=9).value = gmail_id
        if real_sender:
            ws.cell(row=row, column=2).value = real_sender
        filled.append(f"row {row}")

wb.save("eval/label_eval_set.xlsx")
print(f"Filled {len(filled)} rows: {filled}")
