"""Adds 3 new rows to eval/label_eval_set.xlsx testing the narrower
Undecided-vs-Needs-reply ambiguity rule just added to SKILL.md. Expected
Label is recorded here, BEFORE these are ever dumped/labelled -- genuinely
blind: fresh message_ids, no prior label_applied to anchor a re-read on.

Run once: python eval/add_blind_test_rows.py
"""

from openpyxl import load_workbook

ROWS = [
    # sender, subject, body, expected_label
    (6, "Thoughts?", "Thoughts?", "1. Undecided"),
    (6, "Thoughts on the attached pricing proposal?",
     "Thoughts on the attached pricing proposal?", "1. Needs reply"),
    (6, "Can you approve this by 3 PM?",
     "Can you approve this by 3 PM?", "1. Needs reply: ASAP"),
]

wb = load_workbook("eval/label_eval_set.xlsx")
ws = wb["Eval Set"]

start_row = ws.max_row + 1
for i, (sender, subject, body, expected) in enumerate(ROWS):
    row = start_row + i
    ws.cell(row=row, column=1).value = sender
    ws.cell(row=row, column=2).value = "ambiguity-test6@example.com"
    ws.cell(row=row, column=3).value = i + 1
    ws.cell(row=row, column=4).value = "Ambiguity-rule blind test"
    ws.cell(row=row, column=5).value = subject
    ws.cell(row=row, column=6).value = body
    ws.cell(row=row, column=7).value = expected
    ws.cell(row=row, column=11).value = f'=IF(OR(G{row}="",J{row}=""),"",IF(G{row}=J{row},"MATCH","MISMATCH"))'

wb.save("eval/label_eval_set.xlsx")
print(f"Added {len(ROWS)} rows starting at row {start_row}")
